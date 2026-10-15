# Copyright (c) 2025-2026 Patrick Hall, jphall@gwu.edu
# SPDX-License-Identifier: MIT

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import ResearchConfig
from .models import CourseRecord, RunUsage


HEADER_FILL = PatternFill("solid", fgColor="003B5C")
HEADER_FONT = Font(color="FFFFFF", bold=True)
STATUS_FILLS = {
    "explicit": PatternFill("solid", fgColor="C6EFCE"),
    "inferred": PatternFill("solid", fgColor="FFEB9C"),
    "not_evidenced": PatternFill("solid", fgColor="E7E6E6"),
}


def _write_table(sheet, headers: list[str], rows: Iterable[Iterable[Any]]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _set_widths(sheet, widths: dict[str, float]) -> None:
    for name, width in widths.items():
        index = list(sheet.iter_rows(min_row=1, max_row=1))[0]
        for cell in index:
            if cell.value == name:
                sheet.column_dimensions[get_column_letter(cell.column)].width = width


def _hyperlink_column(sheet, header: str) -> None:
    header_map = {cell.value: cell.column for cell in sheet[1]}
    column = header_map.get(header)
    if not column:
        return
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row, column)
        if isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"


def export_workbook(
    path: Path,
    *,
    courses: list[CourseRecord],
    gaps: list[dict[str, Any]],
    config: ResearchConfig,
    usage: RunUsage,
    run_info: dict[str, Any],
) -> None:
    workbook = Workbook()
    selected_sheet = workbook.active
    selected_sheet.title = "Selected courses"
    selected = [course for course in courses if course.selection_status == "selected"]
    _write_table(
        selected_sheet,
        [
            "Course ID", "Provider", "Course", "Instructor", "Audience", "Prerequisites",
            "Provider type", "Date evidence", "Date", "Date source location", "Primary source",
            "All source links", "Selection rationale", "Override outside criteria", "Rigor evidence",
            "Audience-fit evidence", "Ordered outline",
        ],
        (
            [
                course.course_id, course.provider, course.title, course.instructor, course.audience,
                course.prerequisites, course.provider_type, course.date_evidence, course.date_iso,
                course.date_source_location, course.canonical_url, "\n".join(course.source_urls),
                course.selection_rationale, "Yes" if course.override_outside_criteria else "No",
                course.rigor_evidence, course.audience_fit_evidence,
                "\n".join(f"{index}. {item}" for index, item in enumerate(course.outline, 1)),
            ]
            for course in selected
        ),
    )
    _set_widths(selected_sheet, {"Provider": 24, "Course": 34, "Audience": 35, "Primary source": 45,
                                  "Selection rationale": 45, "Ordered outline": 80, "Rigor evidence": 45})
    _hyperlink_column(selected_sheet, "Primary source")

    matrix = workbook.create_sheet("Topic coverage")
    topic_labels = {topic["id"]: topic["label"] for topic in config.topics}
    _write_table(
        matrix,
        ["Course ID", "Provider", "Course", "Topic ID", "Topic", "Coverage status", "Supporting excerpt",
         "Source URL", "Source location", "Explanation", "Responsible use covered elsewhere"],
        (
            [
                course.course_id, course.provider, course.title, evidence.topic_id,
                topic_labels.get(evidence.topic_id, evidence.topic_id), evidence.status, evidence.excerpt,
                evidence.source_url, evidence.source_location, evidence.explanation,
                "Yes" if next((topic.get("covered_elsewhere") for topic in config.topics
                               if topic["id"] == evidence.topic_id), False) else "No",
            ]
            for course in selected
            for evidence in course.coverage
        ),
    )
    status_column = next(cell.column for cell in matrix[1] if cell.value == "Coverage status")
    for row in range(2, matrix.max_row + 1):
        cell = matrix.cell(row, status_column)
        if cell.value in STATUS_FILLS:
            cell.fill = STATUS_FILLS[cell.value]
    _set_widths(matrix, {"Course": 32, "Topic": 35, "Supporting excerpt": 55, "Source URL": 45,
                         "Source location": 28, "Explanation": 50})
    _hyperlink_column(matrix, "Source URL")

    gaps_sheet = workbook.create_sheet("Potential curriculum gaps")
    _write_table(
        gaps_sheet,
        ["Additional topic", "Evidence status", "Supporting courses", "Courses", "Prevalence",
         "Evidence from course materials", "Recommendation / audience relevance", "Interpretation note"],
        (
            [
                gap["topic"], gap["evidence_status"], "\n".join(gap["supporting_courses"]),
                gap["course_count"], gap["prevalence"], "\n".join(gap["evidence"]),
                "\n".join(gap["recommendation"]), gap["note"],
            ]
            for gap in gaps
        ),
    )
    prevalence_column = next(cell.column for cell in gaps_sheet[1] if cell.value == "Prevalence")
    for row in range(2, gaps_sheet.max_row + 1):
        gaps_sheet.cell(row, prevalence_column).number_format = "0.0%"
    _set_widths(gaps_sheet, {"Additional topic": 30, "Supporting courses": 45,
                             "Evidence from course materials": 60,
                             "Recommendation / audience relevance": 60, "Interpretation note": 45})

    unresolved = workbook.create_sheet("Excluded and unresolved")
    _write_table(
        unresolved,
        ["Course ID", "Status", "Provider", "Course", "Primary source", "Reason / uncertainty",
         "Retrieval problems", "Date evidence", "Date", "Duplicate of", "Inclusion override"],
        (
            [
                course.course_id, course.selection_status, course.provider, course.title,
                course.canonical_url, "\n".join(filter(None, [course.exclusion_reason, *course.uncertainties])),
                "\n".join(f"{source.url}: {source.status} — {source.error}" for source in course.sources
                           if source.status != "retrieved"),
                course.date_evidence, course.date_iso, course.duplicate_of,
                "Yes" if course.inclusion_override else "No",
            ]
            for course in courses
            if course.selection_status != "selected"
        ),
    )
    _set_widths(unresolved, {"Provider": 24, "Course": 34, "Primary source": 45,
                             "Reason / uncertainty": 55, "Retrieval problems": 65})
    _hyperlink_column(unresolved, "Primary source")

    info_sheet = workbook.create_sheet("Run information")
    info_rows = [
        ["Profile", config.data["profile_name"]],
        ["Profile path", str(config.path)],
        ["As-of date", config.as_of_date.isoformat()],
        ["Recency months", config.data["recency_months"]],
        ["Target selected courses", f"{config.data['selection_target_min']}–{config.data['selection_target_max']}"],
        ["Selected courses", run_info.get("selected_count", 0)],
        ["Selection shortfall", run_info.get("shortfall", 0)],
        ["Rolling-window cutoff", run_info.get("cutoff_date", "")],
        ["Discovery model", config.model("discovery")],
        ["Analysis model", config.model("analysis")],
        ["Reasoning effort", config.reasoning_effort or "model default"],
        ["API calls", usage.api_calls],
        ["Input tokens", usage.input_tokens],
        ["Output tokens (includes reasoning)", usage.output_tokens],
        ["Reasoning tokens", usage.reasoning_tokens],
        ["Estimated cost", (
            f"${usage.estimated_cost_usd:.4f} from profile-supplied rates; verify against the API dashboard."
            if usage.estimated_cost_usd is not None
            else "Not calculated; supply current rates in the profile and consult the API dashboard."
        )],
        ["Criteria", str(config.data["selection"])],
        ["Seed sources", "\n".join(f"{item['name']}: {item['url']}" for item in config.data["seed_sources"])],
        ["Limitations", "\n".join(run_info.get("limitations", []))],
        ["Evidence semantics", "not_evidenced means the available outline did not establish coverage; it does not mean absent."],
        ["Responsible use", "Configured as covered in another class; shown in the matrix but not proposed as a gap."],
    ]
    _write_table(info_sheet, ["Field", "Value"], info_rows)
    _set_widths(info_sheet, {"Field": 34, "Value": 110})

    for sheet in workbook.worksheets:
        sheet.row_dimensions[1].height = 32
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
