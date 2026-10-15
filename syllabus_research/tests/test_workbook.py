# Copyright (c) 2025-2026 Patrick Hall, jphall@gwu.edu
# SPDX-License-Identifier: MIT

from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook

from syllabus_research.config import ResearchConfig, load_config
from syllabus_research.models import CourseRecord, CoverageEvidence, RunUsage
from syllabus_research.workbook import export_workbook


PROFILE = Path(__file__).parents[1] / "config" / "profiles" / "llm_business_analytics.json"


def test_workbook_has_required_filterable_sheets_and_evidence(tmp_path):
    base = load_config(PROFILE)
    data = deepcopy(base.data)
    data["as_of_date"] = "2026-10-15"
    config = ResearchConfig(base.path, data)
    course = CourseRecord(
        "course-1", "Applied LLMs", provider="Example University", instructor="Ada Example",
        canonical_url="https://example.edu/llms", source_urls=["https://example.edu/llms"],
        date_evidence="Fall 2026", date_iso="2026-09", selection_status="selected",
        selection_rationale="Verified public outline.", outline=["Transformers", "RAG lab"],
        coverage=[CoverageEvidence("rag", "explicit", "RAG lab", "https://example.edu/llms", "L42")],
    )
    output = tmp_path / "comparison.xlsx"
    export_workbook(
        output, courses=[course], gaps=[], config=config, usage=RunUsage(api_calls=2),
        run_info={"selected_count": 1, "shortfall": 19, "cutoff_date": "2025-10-15",
                  "limitations": ["Fixture run"]},
    )
    workbook = load_workbook(output)
    assert workbook.sheetnames == [
        "Selected courses", "Topic coverage", "Potential curriculum gaps",
        "Excluded and unresolved", "Run information",
    ]
    assert workbook["Selected courses"].auto_filter.ref
    assert workbook["Topic coverage"]["G2"].value == "RAG lab"
    assert workbook["Topic coverage"]["H2"].hyperlink.target == "https://example.edu/llms"
